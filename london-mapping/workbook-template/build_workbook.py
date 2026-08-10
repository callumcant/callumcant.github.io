"""
Builds the OneDrive-ready Excel workbook the app reads from, as plain Excel
Tables (no formulas).

    python3 build_workbook.py
    -> London_Project_Mapping_template.xlsx

Column order comes from ../data-dictionary/dictionary.json, which is also what
js/data/table-schemas.js is generated from. Reading it directly here means the
workbook and the app cannot disagree about column order — and they must not,
because the app maps Graph's positional row `values` arrays onto field names by
that order.

To change the model: edit dictionary.json, run
data-dictionary/generate_schemas.py, then run this.

Every row written here is a clearly-fake EXAMPLE. Delete them before pasting in
real data, and never commit real member/employer/dispute data to this repo.
"""
import json
import pathlib

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = pathlib.Path(__file__).parent
DICTIONARY = HERE.parent / "data-dictionary" / "dictionary.json"
OUTPUT = HERE / "London_Project_Mapping_template.xlsx"

FONT = "Arial"

LONDON_BOROUGHS = [
    "Barking and Dagenham", "Barnet", "Bexley", "Brent", "Bromley", "Camden",
    "Croydon", "Ealing", "Enfield", "Greenwich", "Hackney",
    "Hammersmith and Fulham", "Haringey", "Harrow", "Havering", "Hillingdon",
    "Hounslow", "Islington", "Kensington and Chelsea", "Kingston Upon Thames",
    "Lambeth", "Lewisham", "Merton", "Newham", "Redbridge",
    "Richmond Upon Thames", "Southwark", "Sutton", "Tower Hamlets (&CoL)",
    "Waltham Forest", "Wandsworth", "Westminster",
]

# Confirmed for 2026/27. Greenwich was a project branch earlier and is not one
# now. Changing this is a Yes/No edit in the BranchFacts sheet, not a code change.
PROJECT_BRANCHES = {
    "Havering", "Hillingdon", "Bromley", "Wandsworth", "Kensington and Chelsea",
}
PROJECT_MATS = [
    "Oscar Romero", "COLA", "Haberdashers", "Compass Eko", "Orchard Hill",
]

# Compass Eko is a merger of two trusts, so GIAS and Stratum may still name
# either half. Seeding the aliases means it appears as one MAT from day one
# rather than silently splitting in two.
MAT_ALIAS_SEEDS = [
    ["Compass Partnership of Schools", "Compass Eko"],
    ["EKO Trust", "Compass Eko"],
]


def rows_for(table):
    """Per-table seed rows. Everything else gets its single EXAMPLE row."""
    name = table["name"]
    if name == "BranchFacts":
        return [[b, "Yes" if b in PROJECT_BRANCHES else "No", 0] for b in LONDON_BOROUGHS]
    if name == "MatFacts":
        return [[m, "Yes", "No"] for m in PROJECT_MATS]
    if name == "MatAliases":
        return MAT_ALIAS_SEEDS
    example = table.get("example")
    return [example] if example else []


def build():
    doc = json.loads(DICTIONARY.read_text())
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for table in doc["tables"]:
        ws = wb.create_sheet(table["name"])
        headers = [f["header"] for f in table["fields"]]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(name=FONT, bold=True)

        body = rows_for(table)
        for row in body:
            ws.append(row)

        n_rows = len(body) + 1
        n_cols = len(headers)
        last_col = openpyxl.utils.get_column_letter(n_cols)
        excel_table = Table(displayName=table["name"], ref=f"A1:{last_col}{n_rows}")
        excel_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(excel_table)

        for idx, header in enumerate(headers, 1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = max(12, len(header) + 2)
        for row in ws.iter_rows(min_row=2, max_row=n_rows, max_col=n_cols):
            for cell in row:
                cell.font = Font(name=FONT)

        for col_letter, options in table.get("dv", {}).items():
            dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
            dv.add(f"{col_letter}2:{col_letter}2000")
            ws.add_data_validation(dv)

        note = ws.cell(row=n_rows + 2, column=1, value=f"Note: {table['note']}"
                       if table.get("note") else f"Note: {table['description']}")
        note.font = Font(name=FONT, italic=True, size=9, color="898781")

        source_note = ws.cell(
            row=n_rows + 3, column=1,
            value="Column order must match data-dictionary/dictionary.json — the app reads columns by position.",
        )
        source_note.font = Font(name=FONT, italic=True, size=9, color="898781")

    wb.save(OUTPUT)
    total = sum(len(t["fields"]) for t in doc["tables"])
    print(f"Wrote {OUTPUT.name} — {len(doc['tables'])} tables, {total} columns")


if __name__ == "__main__":
    build()
